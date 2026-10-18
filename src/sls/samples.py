"""Synthetic sample data for tests and demonstrations.

Everything here is fabricated: coordinates, rule values and license records are NOT
official data and must never be used for real licensing decisions.
"""
from __future__ import annotations

from pathlib import Path

import ezdxf
from openpyxl import Workbook

# Synthetic origin in JTM (EPSG:3066) range, exercising large coordinates and the CRS checks.
OX, OY = 353_000.0, 355_500.0  # synthetic geometry placed at JTM-like coordinates near Wadi Musa

SAMPLE_PARCEL = [(0.0, 0.0), (32.0, 0.0), (30.0, 26.0), (1.5, 25.0)]  # front edge = first edge (road side)
SAMPLE_GROUND_FLOOR = [(2.0, 5.0), (22.0, 5.0), (22.0, 18.0), (2.0, 18.0)]  # intrudes into left side setback
SAMPLE_FIRST_FLOOR = [(4.0, 6.0), (20.0, 6.0), (20.0, 16.0), (4.0, 16.0)]
SAMPLE_ANNEX = [(24.0, 14.0), (28.0, 14.0), (28.0, 21.0), (24.0, 21.0)]


def _g(pts: list[tuple[float, float]]) -> list[tuple[float, float]]:
    return [(OX + x, OY + y) for x, y in pts]


def write_sample_dxf(path: Path, units_code: int = 6, include_unsupported: bool = True) -> Path:
    doc = ezdxf.new("R2018", setup=True)
    doc.header["$INSUNITS"] = units_code
    for name, color in [("PARCEL", 4), ("BLDG_GF", 2), ("BLDG_1F", 30), ("WALL", 7), ("FENCE", 34),
                        ("ROAD", 8), ("NEIGHBOR", 3), ("TEXT", 7), ("WELL", 5), ("HATCH", 252)]:
        doc.layers.add(name, color=color)
    msp = doc.modelspace()
    msp.add_lwpolyline(_g(SAMPLE_PARCEL), close=True, dxfattribs={"layer": "PARCEL"})
    msp.add_lwpolyline(_g(SAMPLE_GROUND_FLOOR), close=True, dxfattribs={"layer": "BLDG_GF"})
    msp.add_lwpolyline(_g(SAMPLE_ANNEX), close=True, dxfattribs={"layer": "BLDG_GF"})
    msp.add_lwpolyline(_g(SAMPLE_FIRST_FLOOR), close=True, dxfattribs={"layer": "BLDG_1F"})
    # Internal walls (lines) and boundary fence (open polyline along three sides).
    msp.add_line(*_g([(12.0, 5.0), (12.0, 18.0)]), dxfattribs={"layer": "WALL"})
    msp.add_line(*_g([(2.0, 11.0), (12.0, 11.0)]), dxfattribs={"layer": "WALL"})
    msp.add_lwpolyline(_g([(0.0, 0.0), (1.5, 25.0), (30.0, 26.0), (32.0, 0.0)]), dxfattribs={"layer": "FENCE"})
    # Road in front and neighbouring parcels.
    msp.add_lwpolyline(_g([(-10.0, -8.0), (45.0, -8.0), (45.0, -0.0), (-10.0, 0.0)]), close=True,
                       dxfattribs={"layer": "ROAD"})
    msp.add_lwpolyline(_g([(32.0, 0.0), (55.0, 0.0), (52.0, 27.0), (30.0, 26.0)]), close=True,
                       dxfattribs={"layer": "NEIGHBOR"})
    # Arabic and English labels.
    msp.add_text("قطعة تجريبية رقم ١٢٣", height=0.8, dxfattribs={"layer": "TEXT"}).set_placement(_g([(8.0, 22.0)])[0])
    msp.add_text("SYNTHETIC SURVEY - NOT OFFICIAL", height=0.6,
                 dxfattribs={"layer": "TEXT"}).set_placement(_g([(4.0, -4.0)])[0])
    msp.add_mtext("شارع تجريبي", dxfattribs={"layer": "TEXT", "char_height": 0.8,
                                              "insert": _g([(15.0, -4.5)])[0]})
    # A well drawn as a block reference.
    blk = doc.blocks.new("WELL_SYMBOL")
    blk.add_circle((0, 0), 0.6)
    blk.add_line((-0.6, 0), (0.6, 0))
    msp.add_blockref("WELL_SYMBOL", _g([(26.0, 6.0)])[0], dxfattribs={"layer": "WELL"})
    msp.add_arc(_g([(26.0, 6.0)])[0], 1.2, 0, 180, dxfattribs={"layer": "WELL"})
    if include_unsupported:
        hatch = msp.add_hatch(color=252, dxfattribs={"layer": "HATCH"})
        hatch.paths.add_polyline_path(_g([(24.0, 14.0), (28.0, 14.0), (28.0, 21.0), (24.0, 21.0)]), is_closed=True)
        msp.add_spline(_g([(5.0, 20.0), (8.0, 21.0), (11.0, 20.0), (14.0, 21.5)]), dxfattribs={"layer": "TEXT"})
    doc.saveas(str(path))
    return path


RULE_HEADERS = ["رمز التصنيف", "التصنيف", "Classification", "أمامي (م)", "جانبي (م)", "خلفي (م)",
                "نسبة البناء القصوى", "عدد الطوابق", "معامل الطوابق", "ملاحظات"]

SAMPLE_RULE_ROWS = [
    ["SYN-A", "سكن تجريبي أ", "Synthetic residential A", 5, 4, 4, 0.45, 4, 1.8, "قيم اصطناعية للاختبار فقط"],
    ["SYN-B", "سكن تجريبي ب", "Synthetic residential B", 4, 3, 3, 0.55, 4, 2.2, "قيم اصطناعية للاختبار فقط"],
    ["SYN-C", "سكن تجريبي ج", "Synthetic residential C", 3, 2.5, 2.5, 0.60, 3, 1.8, "قيم اصطناعية للاختبار فقط"],
    ["SYN-COM", "تجاري تجريبي", "Synthetic commercial", 0, 0, 3, 0.80, 5, 4.0, "قيم اصطناعية للاختبار فقط"],
]

SAMPLE_RULE_MAPPING = {
    "columns": {"code": "رمز التصنيف", "name_ar": "التصنيف", "name_en": "Classification",
                "front": "أمامي (م)", "side": "جانبي (م)", "rear": "خلفي (م)",
                "max_coverage_pct": "نسبة البناء القصوى", "max_floors": "عدد الطوابق",
                "max_far": "معامل الطوابق", "notes": "ملاحظات"},
    "extra_setbacks": {},
}


def write_sample_rules_xlsx(path: Path, rows: list[list[object]] | None = None) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Setbacks"
    ws.append(["جدول ارتدادات اصطناعي — SYNTHETIC RULE TABLE, NOT OFFICIAL"])
    ws.append(RULE_HEADERS)
    for row in rows if rows is not None else SAMPLE_RULE_ROWS:
        ws.append(row)
        ws.cell(row=ws.max_row, column=7).number_format = "0%"
    wb.create_sheet("Notes").append(["Values are fabricated for testing."])
    wb.save(str(path))
    return path


SAMPLE_LAYER_ROLES = {
    "PARCEL": ("property_boundary", ""),
    "BLDG_GF": ("building", "G"),
    "BLDG_1F": ("building", "1"),
    "WALL": ("wall", ""),
    "FENCE": ("fence", ""),
    "ROAD": ("road", ""),
    "NEIGHBOR": ("neighbor_parcel", ""),
    "TEXT": ("annotation", ""),
    "WELL": ("well", ""),
}


def assign_sample_roles(drawing) -> str:  # type: ignore[no-untyped-def]
    """Assign layer roles for the synthetic survey; returns the parcel boundary entity id."""
    for name, (role, floor) in SAMPLE_LAYER_ROLES.items():
        if drawing.has_layer(name):
            layer = drawing.layer(name)
            layer.role = role
            layer.floor_key = floor
    return next(e.id for e in drawing.entities.values() if e.layer == "PARCEL")


SAMPLE_EDGE_TYPES = ["front", "side", "rear", "side"]

"""Land-use setback rule tables from Excel workbooks (.xlsx/.xlsm via openpyxl, legacy .xls via xlrd).

Workflow: inspect -> preview sheet -> explicit column mapping -> validate -> import.

Real rule tables mix plain numbers with conditional text such as "صفر - 2", "4م بعد 14م" or
"حسب المخطط". Nothing is inferred: every value keeps its exact source text, and a number is
stored only when the text is an unambiguous plain value ("5", "2.5", "صفر", "39%", "4م").
Rows whose setbacks are conditional are imported and flagged; calculations for them stop
with the source text until an authorised reviewer records the applicable values for a case.
Structural problems (missing or duplicate codes, negative numbers, unmapped columns) block
the import.
"""
from __future__ import annotations

import math
import re
from collections.abc import Mapping, Sequence
from dataclasses import asdict, dataclass, field, fields
from pathlib import Path
from typing import Any

from ..core.errors import AmbiguousRuleError, MissingRuleError, RuleImportError

REQUIRED_FIELDS = ("code", "front", "side", "rear")
OPTIONAL_FIELDS = ("name_ar", "name_en", "max_coverage_pct", "max_floors", "max_far", "notes")
BASE_BOUNDARY_TYPES = ("front", "side", "rear")
RULE_STATUSES = ("numeric", "conditional", "no_setbacks")

_ARABIC_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩٫٬", "0123456789.,")
_ZERO_WORDS = {"صفر"}
_NUM = r"[0-9]+(?:\.[0-9]+)?"
_PLAIN = re.compile(rf"^{_NUM}$")
_METRES = re.compile(rf"^({_NUM})\s*(?:م|m)$")
_PERCENT = re.compile(rf"^(?:%\s*({_NUM})|({_NUM})\s*%)$")

# Header names recognised for the automatic mapping *suggestion* (the user confirms it).
_SUGGEST = {
    "code": ("zone_id", "code", "zone", "رمز", "رمز التصنيف", "رقم الصفة"),
    "name_ar": ("name", "name_ar", "الاسم", "التصنيف", "الصفة", "صفة الاستعمال"),
    "name_en": ("name_en", "classification", "english name"),
    "front": ("setback_front", "front", "أمامي", "امامي", "أمامي (م)", "الارتداد الامامي"),
    "side": ("setback_side", "side", "جانبي", "جانبي (م)", "الارتداد الجانبي"),
    "rear": ("setback_back", "setback_rear", "rear", "back", "خلفي", "خلفي (م)", "الارتداد الخلفي"),
    "max_coverage_pct": ("built_up_area_percent", "coverage", "نسبة البناء", "نسبة البناء القصوى"),
    "max_floors": ("floor_no", "floors", "max_floors", "عدد الطوابق"),
    "max_far": ("floor_area_ratio", "far", "معامل الطوابق"),
    "notes": ("remark", "notes", "ملاحظات"),
}


@dataclass
class Cell:
    value: Any
    number_format: str = "General"


@dataclass
class SheetPreview:
    sheet: str
    header_row: int  # 1-based
    headers: list[str]
    rows: list[list[Cell]]  # data rows after the header
    first_data_row: int


@dataclass
class ColumnMapping:
    """Maps logical fields to header names in the sheet.

    ``extra_setbacks`` maps additional boundary types (e.g. "corner", "side_2") to headers,
    for rule tables that define more than front/side/rear.
    """

    columns: dict[str, str] = field(default_factory=dict)
    extra_setbacks: dict[str, str] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)

    @classmethod
    def from_dict(cls, d: Mapping[str, Any]) -> "ColumnMapping":
        return cls(dict(d.get("columns", {})), dict(d.get("extra_setbacks", {})))


@dataclass
class LandUseRule:
    code: str
    setbacks: dict[str, float]  # only unambiguous numeric values
    source_row: int
    name_ar: str = ""
    name_en: str = ""
    max_coverage_pct: float | None = None
    max_floors: int | None = None
    max_far: float | None = None
    notes: str = ""
    status: str = "numeric"
    setbacks_raw: dict[str, str] = field(default_factory=dict)  # exact source text per boundary type
    coverage_raw: str = ""
    floors_raw: str = ""
    far_raw: str = ""
    warnings: list[str] = field(default_factory=list)
    source: dict[str, str] = field(default_factory=dict)  # every column of the source row, verbatim

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)

    @classmethod
    def from_dict(cls, d: Mapping[str, Any]) -> "LandUseRule":
        known = {f.name for f in fields(cls)}
        return cls(**{k: v for k, v in d.items() if k in known})  # older snapshots lack new fields

    def label(self, lang: str) -> str:
        name = self.name_ar if lang == "ar" else (self.name_en or self.name_ar)
        return f"{self.code} — {name}" if name else self.code

    def raw_setback(self, btype: str) -> str:
        return self.setbacks_raw.get(btype, "")


@dataclass
class RowError:
    row: int
    column: str
    message: str


@dataclass
class RuleValidation:
    rules: list[LandUseRule]
    errors: list[RowError]
    skipped_blank_rows: int
    warnings: list[RowError] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors and bool(self.rules)

    def count(self, status: str) -> int:
        return sum(1 for r in self.rules if r.status == status)


# ------------------------------------------------------------------------- reading


def _is_xls(path: Path) -> bool:
    return Path(path).suffix.lower() == ".xls"


def list_sheets(path: Path) -> list[str]:
    if _is_xls(path):
        import xlrd

        try:
            book = xlrd.open_workbook(str(path), on_demand=True)
        except xlrd.XLRDError as exc:
            raise RuleImportError(f"Cannot read {Path(path).name}: {exc}") from exc
        try:
            return book.sheet_names()
        finally:
            book.release_resources()
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _read_xls(path: Path, sheet: str) -> list[list[Cell]]:
    import xlrd

    try:
        book = xlrd.open_workbook(str(path), formatting_info=True)
    except xlrd.XLRDError as exc:
        raise RuleImportError(f"Cannot read {Path(path).name}: {exc}") from exc
    if sheet not in book.sheet_names():
        raise RuleImportError(f"Worksheet '{sheet}' not found in {Path(path).name}.")
    sh = book.sheet_by_name(sheet)
    rows: list[list[Cell]] = []
    for r in range(sh.nrows):
        row: list[Cell] = []
        for c in range(sh.ncols):
            cell = sh.cell(r, c)
            fmt = "General"
            try:
                fmt = book.format_map[book.xf_list[cell.xf_index].format_key].format_str
            except (IndexError, KeyError, AttributeError):
                pass
            if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                value: Any = None
            elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                value = bool(cell.value)
            elif cell.ctype == xlrd.XL_CELL_ERROR:
                value = None
            else:
                value = cell.value
            row.append(Cell(value, fmt))
        rows.append(row)
    return rows


def _read_xlsx(path: Path, sheet: str, max_rows: int) -> list[list[Cell]]:
    from openpyxl import load_workbook

    # data_only=True: use cached formula results. Formula cells without cached values
    # (workbook never saved by Excel) come back as None and are flagged as missing.
    wb = load_workbook(path, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise RuleImportError(f"Worksheet '{sheet}' not found in {Path(path).name}.")
        ws = wb[sheet]
        return [[Cell(c.value, c.number_format or "General") for c in r]
                for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row, max_rows))]
    finally:
        wb.close()


def read_sheet(path: Path, sheet: str, header_row: int = 1, max_rows: int = 5000) -> SheetPreview:
    all_rows = _read_xls(path, sheet) if _is_xls(path) else _read_xlsx(path, sheet, max_rows + header_row)
    if header_row < 1 or header_row > len(all_rows):
        raise RuleImportError(f"Header row {header_row} is outside the sheet (1..{len(all_rows)}).")
    headers = [_text(c.value) for c in all_rows[header_row - 1]]
    while headers and not headers[-1]:
        headers.pop()
    if not any(headers):
        raise RuleImportError(f"Header row {header_row} is empty.")
    dupes = {h for h in headers if h and headers.count(h) > 1}
    if dupes:
        raise RuleImportError(f"Duplicate column headers: {', '.join(sorted(dupes))}")
    rows = [r[: len(headers)] + [Cell(None)] * (len(headers) - len(r))
            for r in all_rows[header_row: header_row + max_rows]]
    return SheetPreview(sheet, header_row, headers, rows, header_row + 1)


def suggest_mapping(headers: Sequence[str]) -> ColumnMapping:
    """Suggest a mapping from recognised header names. Only a suggestion: the user confirms it."""
    norm = {h.strip().casefold(): h for h in headers if h}
    columns: dict[str, str] = {}
    for fld, names in _SUGGEST.items():
        for n in names:
            if n.casefold() in norm and norm[n.casefold()] not in columns.values():
                columns[fld] = norm[n.casefold()]
                break
    return ColumnMapping(columns, {})


# ------------------------------------------------------------------------- parsing


def _text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return " ".join(str(v).split())  # collapse whitespace and line breaks


@dataclass(frozen=True)
class Parsed:
    raw: str
    number: float | None
    state: str  # "blank" | "number" | "text"


def parse_value(cell: Cell | None, kind: str = "length") -> Parsed:
    """Parse one rule value. ``kind`` is "length", "percent" or "count".

    Only unambiguous forms become numbers; anything else stays text (state "text").
    """
    if cell is None or cell.value is None:
        return Parsed("", None, "blank")
    v = cell.value
    if isinstance(v, bool):
        return Parsed(str(v), None, "text")
    if isinstance(v, (int, float)):
        num = float(v)
        raw = _text(v)
        if not math.isfinite(num):
            return Parsed(raw, None, "text")
        if kind == "percent":
            if "%" in (cell.number_format or ""):
                num *= 100.0  # Excel stores 60% as 0.6
                raw = f"{num:g}%"
            elif 0 < num <= 1:
                return Parsed(raw, None, "text")  # 0.45 without a % format: 45% or 0.45%? ambiguous
        if kind == "count" and not num.is_integer():
            return Parsed(raw, None, "text")
        return Parsed(raw, num, "number")
    raw = _text(v)
    s = raw.translate(_ARABIC_DIGITS).replace(" ", "")
    if not s:
        return Parsed("", None, "blank")
    if s in _ZERO_WORDS:
        return Parsed(raw, 0.0, "number")
    if kind == "percent":
        m = _PERCENT.match(s)
        if m:
            return Parsed(raw, float(m.group(1) or m.group(2)), "number")
        if _PLAIN.match(s) and float(s) == 0:
            return Parsed(raw, 0.0, "number")
        if _PLAIN.match(s) and float(s) > 1:
            return Parsed(raw, float(s), "number")
        return Parsed(raw, None, "text")
    if _PLAIN.match(s):
        num = float(s)
        if kind == "count" and not num.is_integer():
            return Parsed(raw, None, "text")
        return Parsed(raw, num, "number")
    if kind == "length":
        m = _METRES.match(s)
        if m:
            return Parsed(raw, float(m.group(1)), "number")
    return Parsed(raw, None, "text")


def validate(preview: SheetPreview, mapping: ColumnMapping) -> RuleValidation:
    errors: list[RowError] = []
    warnings: list[RowError] = []
    header_index = {h: i for i, h in enumerate(preview.headers) if h}
    for fld in REQUIRED_FIELDS:
        if not mapping.columns.get(fld):
            errors.append(RowError(preview.header_row, fld, f"Required field '{fld}' is not mapped."))
    all_cols = {**mapping.columns, **{f"setback:{k}": v for k, v in mapping.extra_setbacks.items()}}
    for fld, header in all_cols.items():
        if header and header not in header_index:
            errors.append(RowError(preview.header_row, fld, f"Mapped column '{header}' does not exist."))
    for key in mapping.extra_setbacks:
        if key in BASE_BOUNDARY_TYPES or not key.strip():
            errors.append(RowError(preview.header_row, key, f"Invalid extra boundary type '{key}'."))
    if errors:
        return RuleValidation([], errors, 0)

    def cell(row: list[Cell], header: str | None) -> Cell | None:
        if not header:
            return None
        idx = header_index[header]
        return row[idx] if idx < len(row) else None

    rules: list[LandUseRule] = []
    seen: dict[str, int] = {}
    blank = 0
    sb_cols = {k: mapping.columns[k] for k in BASE_BOUNDARY_TYPES}
    sb_cols.update(mapping.extra_setbacks)
    for offset, row in enumerate(preview.rows):
        rownum = preview.first_data_row + offset
        if all(c.value is None or (isinstance(c.value, str) and not c.value.strip()) for c in row):
            blank += 1
            continue
        code = _text(cell(row, mapping.columns["code"]).value if cell(row, mapping.columns["code"]) else None)
        if not code:
            errors.append(RowError(rownum, "code", "Classification code is missing."))
            continue
        if code in seen:
            errors.append(RowError(rownum, "code", f"Duplicate classification code '{code}' (also row {seen[code]})."))
            continue
        seen[code] = rownum
        row_ok = True
        row_warn: list[str] = []
        setbacks: dict[str, float] = {}
        setbacks_raw: dict[str, str] = {}
        text_types: list[str] = []
        blank_base: list[str] = []
        for btype, header in sb_cols.items():
            p = parse_value(cell(row, header), "length")
            setbacks_raw[btype] = p.raw
            if p.state == "blank":
                if btype in BASE_BOUNDARY_TYPES:
                    blank_base.append(btype)
                continue
            if p.state == "text":
                text_types.append(btype)
                continue
            if p.number is not None and p.number < 0:
                errors.append(RowError(rownum, header, f"Setback '{btype}' is negative ({p.raw})."))
                row_ok = False
                continue
            setbacks[btype] = float(p.number)  # type: ignore[arg-type]
        if len(blank_base) == len(BASE_BOUNDARY_TYPES) and not text_types:
            status = "no_setbacks"
            row_warn.append("No setback values are defined for this classification.")
        elif text_types or blank_base:
            status = "conditional"
            parts = [f"{t}='{setbacks_raw[t]}'" for t in text_types] + [f"{t}=(blank)" for t in blank_base]
            row_warn.append("Setbacks not plain numbers: " + ", ".join(parts))
        else:
            status = "numeric"

        optional: dict[str, Any] = {}
        raws: dict[str, str] = {}
        for fld, kind, raw_key in (("max_coverage_pct", "percent", "coverage_raw"), ("max_floors", "count", "floors_raw"),
                                   ("max_far", "length", "far_raw")):
            header = mapping.columns.get(fld)
            if not header:
                continue
            p = parse_value(cell(row, header), kind)
            raws[raw_key] = p.raw
            if p.state == "text":
                row_warn.append(f"{fld}: '{p.raw}' is not a plain value; it will not be evaluated automatically.")
                continue
            if p.number is None:
                continue
            if p.number < 0:
                errors.append(RowError(rownum, header, f"'{fld}' is negative ({p.raw})."))
                row_ok = False
                continue
            if fld == "max_coverage_pct" and p.number > 100:
                errors.append(RowError(rownum, header, f"Coverage {p.raw} exceeds 100%."))
                row_ok = False
                continue
            if p.number == 0 and fld in ("max_coverage_pct", "max_floors"):
                row_warn.append(f"{fld} is 0 — confirm whether this means 'not permitted' or 'not specified'.")
            optional[fld] = int(p.number) if fld == "max_floors" else p.number
        if not row_ok:
            continue
        texts = {f: _text(cell(row, mapping.columns.get(f)).value) if cell(row, mapping.columns.get(f)) else ""
                 for f in ("name_ar", "name_en", "notes")}
        source = {h: _text(row[i].value) for h, i in header_index.items() if i < len(row)}
        rules.append(LandUseRule(code=code, setbacks=setbacks, source_row=rownum, status=status,
                                 setbacks_raw=setbacks_raw, warnings=row_warn, source=source,
                                 **texts, **optional, **raws))
        warnings += [RowError(rownum, code, w) for w in row_warn]
    if not rules and not errors:
        errors.append(RowError(preview.first_data_row, "", "The sheet contains no rule rows."))
    return RuleValidation(rules, errors, blank, warnings)


def select_rule(rules: Sequence[LandUseRule], code: str) -> LandUseRule:
    """Return the single rule for ``code``; never falls back to another rule."""
    code = code.strip()
    if not code:
        raise MissingRuleError("No land-use classification has been selected for this case.")
    matches = [r for r in rules if r.code == code]
    if not matches:
        raise MissingRuleError(f"Land-use classification '{code}' is not defined in the applied rule version.",
                               code=code)
    if len(matches) > 1:
        raise AmbiguousRuleError(f"Land-use classification '{code}' is defined more than once.", code=code)
    return matches[0]
